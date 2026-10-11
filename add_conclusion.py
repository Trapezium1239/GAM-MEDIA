from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

doc = Document(r'C:\Users\henry\Desktop\GAM APP\Thesis_Corrected.docx')

# Find the REFERENCES heading
ref_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'REFERENCE':
        ref_idx = i
        break

print(f'References at para {ref_idx}')
ref_heading = doc.paragraphs[ref_idx]._element

# Conclusion Chapter content
conclusion_content = [
    ('CHAPTER SIX', 'Heading 1'),
    ('CONCLUSION AND RECOMMENDATIONS', 'Heading 1'),
    ('6.1 Conclusion', 'Heading 2'),
    ('This study demonstrated that a brief, culturally adapted mental health training programme can significantly improve depression literacy, identification confidence, referral willingness, and reduce stigma among religious educators working with adolescents in Ibadan, Nigeria. The 12-hour training, delivered over two days by mental health professionals, produced large effect sizes (Cohen\'s d = 1.39-2.01) across all outcome measures, with all four hypotheses supported at p < .001.', 'Normal'),
    ('The findings address a critical gap in adolescent mental health services in Nigeria, where the treatment gap exceeds 90% (Okechukwu, 2020) and religious institutions are trusted but underutilised resources for mental health support (Aluh et al., 2018). By equipping religious educators with evidence-based knowledge and practical skills, the training bridges the divide between faith-based community support and formal mental health care.', 'Normal'),
    ('Key conclusions:', 'Normal'),
    ('1. Religious educators can be effective gatekeepers for adolescent mental health. With minimal training, they can accurately identify depression symptoms and make appropriate referrals.', 'Normal'),
    ('2. Culturally sensitive training that respects and integrates faith perspectives increases acceptability and reduces resistance to mental health concepts in religious settings.', 'Normal'),
    ('3. Stigma reduction is achievable even in contexts where mental illness is heavily stigmatised, through myth-busting, contact-based strategies, and reframing mental health as consistent with religious values of compassion.', 'Normal'),
    ('4. The single-group pre-post design limits causal inference. While the magnitude and consistency of effects across domains are encouraging, a randomised controlled trial is required to establish efficacy definitively.', 'Normal'),
    ('6.2 Recommendations', 'Heading 2'),
    ('6.2.1 For Policy and Practice', 'Heading 3'),
    ('1. The Oyo State Ministry of Health should consider integrating mental health training for religious educators into the state\'s mental health action plan, leveraging the existing network of churches across the state.', 'Normal'),
    ('2. The Mental Health Bill (2023) should explicitly recognise the role of trained faith-based community workers in early identification and referral, providing a legal framework for their involvement.', 'Normal'),
    ('3. A standardised national training curriculum for religious educators should be developed, drawing on the modules validated in this study, with accreditation by the Federal Ministry of Health.', 'Normal'),
    ('4. Referral pathways between churches and primary health care centres should be formalised through memoranda of understanding, with designated mental health liaison officers at each facility.', 'Normal'),
    ('5. Funding should be allocated for periodic refresher training (annually) and supervision by mental health professionals to maintain competency and address emerging issues.', 'Normal'),
    ('6.2.2 For Religious Institutions', 'Heading 3'),
    ('1. Church leadership should prioritise mental health literacy as part of pastoral care, designating trained educators as mental health champions within each congregation.', 'Normal'),
    ('2. Churches should establish confidential documentation systems for mental health concerns, protecting adolescent privacy while enabling follow-up.', 'Normal'),
    ('3. Regular mental health awareness sessions should be incorporated into youth programmes, normalising help-seeking and reducing stigma among adolescents directly.', 'Normal'),
    ('4. Inter-faith collaboration on mental health should be encouraged, sharing best practices across Christian, Muslim, and traditional faith communities.', 'Normal'),
    ('6.2.3 For Research', 'Heading 3'),
    ('1. A cluster randomised controlled trial with 20-30 churches per arm is needed to confirm the training\'s efficacy and assess long-term outcomes (6-12 month follow-up).', 'Normal'),
    ('2. Implementation research should explore barriers and facilitators to scaling the programme across diverse Nigerian contexts (urban/rural, North/South, different denominations).', 'Normal'),
    ('3. Cost-effectiveness analysis should compare the training programme to alternative strategies for reducing the adolescent mental health treatment gap.', 'Normal'),
    ('4. Digital adaptations (mobile-based training, tele-supervision) should be developed and tested for cost-effective scaling.', 'Normal'),
    ('5. Research should examine the downstream effects on adolescents themselves: do trained educators\' actions lead to earlier help-seeking, reduced symptom severity, and better functioning?', 'Normal'),
    ('6.3 Final Statement', 'Heading 2'),
    ('Adolescent depression is a silent epidemic in Nigeria, concealed by stigma, poverty, and a fragmented mental health system. Religious educators—present in every community, trusted by families, and invested in young people\'s wellbeing—represent an untapped workforce for early identification and referral. This study shows that with just 12 hours of training, they can become competent, confident, and willing agents of change. The path from evidence to implementation is clear; the moral imperative to act is undeniable.', 'Normal'),
]

# Insert before references
last_element = ref_heading
for text, style_name in reversed(conclusion_content):
    if text == '':
        new_p = doc.add_paragraph()
    else:
        new_p = doc.add_paragraph(text)
        if style_name != 'Normal':
            new_p.style = doc.styles[style_name]
    last_element.addprevious(new_p._element)

doc.save(r'C:\Users\henry\Desktop\GAM APP\Thesis_Corrected.docx')
print('Conclusion chapter added')