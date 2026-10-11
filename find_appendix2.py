from docx import Document

doc = Document(r'C:\Users\henry\Desktop\GAM APP\Thesis_Corrected.docx')

# Find the Appendix and questionnaire items
for i, p in enumerate(doc.paragraphs):
    text = p.text[:150]
    if 'APPENDIX' in text or 'SECTION 1' in text or 'Serial No' in text or 'I can identify' in text or 'Which of these' in text or 'People with mental' in text or 'Adolescents with' in text or 'How would you rate' in text or 'Did you get' in text or 'What I liked' in text or 'DATA TABLE' in text or 'SUMMARY STATISTICS' in text:
        print(f'Para {i}: [{p.style.name}] {text}')
    if i > len(doc.paragraphs):
        break