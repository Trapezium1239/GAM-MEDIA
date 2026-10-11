from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document(r'C:\Users\henry\Desktop\GAM APP\Thesis_Corrected.docx')

# Extract all paragraphs with their style and text
all_paras = []
for p in doc.paragraphs:
    all_paras.append({
        'style': p.style.name,
        'text': p.text,
        'element': p._element
    })

# Define section boundaries by paragraph index
# Based on the Heading 1 structure
sections = {
    'Ch1': (1, 42),      # Chapter One: Introduction (para 1-42)
    'Ch2': (43, 245),    # Chapter Two: Literature Review (para 43-245)
    'Ch3': (119, 245),   # Chapter Three starts at 3.1 Study Design (para 119) - overlaps with Ch2 end
    'Results': (246, 246), # RESULTS heading only
    'Ch4': (358, 401),   # CHAPTER FOUR content (para 358-401)
    'Ch5': (333, 357),   # CHAPTER FIVE content (para 333-357)
    'Ch6': (303, 332),   # CHAPTER SIX content (para 303-332)
    'Ref': (247, 302),   # REFERENCE section (para 247-302)
    'App': (402, len(all_paras)-1),  # APPENDIX (para 402+)
}

# The issue: Ch3 overlaps with Ch2. Let me get the actual Ch3 content
# Ch3 starts at para 119 (3.1 Study Design) and goes until para 245 (before RESULTS)
# But Ch2 goes from 43 to 245. So Ch3 is embedded in Ch2.

# Let me extract content for each chapter properly
def get_content(start_idx, end_idx):
    content = []
    for i in range(start_idx, min(end_idx + 1, len(all_paras))):
        p = all_paras[i]
        if p['text'].strip() or p['style'].startswith('Heading'):
            content.append((p['style'], p['text']))
    return content

# Chapter 1
ch1_content = get_content(1, 42)
# Chapter 2
ch2_content = get_content(43, 118)  # Up to before Ch3
# Chapter 3 (Methodology)
ch3_content = get_content(119, 245)
# Results (Chapter 4)
# The RESULTS heading is at 246, but the actual content is at 358-401 (CHAPTER FOUR)
ch4_content = get_content(358, 401)
# Discussion (Chapter 5)
ch5_content = get_content(333, 357)
# Conclusion (Chapter 6)
ch6_content = get_content(303, 332)
# References
ref_content = get_content(247, 302)
# Appendix
app_content = get_content(402, len(all_paras)-1)

print(f'Ch1: {len(ch1_content)} items')
print(f'Ch2: {len(ch2_content)} items')
print(f'Ch3: {len(ch3_content)} items')
print(f'Ch4: {len(ch4_content)} items')
print(f'Ch5: {len(ch5_content)} items')
print(f'Ch6: {len(ch6_content)} items')
print(f'Ref: {len(ref_content)} items')
print(f'App: {len(app_content)} items')

# Create new document with correct order
new_doc = Document()

# Set styles
style = new_doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for level in range(1, 4):
    h_style = new_doc.styles[f'Heading {level}']
    h_style.font.name = 'Times New Roman'
    h_style.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        h_style.font.size = Pt(16)
        h_style.font.bold = True
        h_style.paragraph_format.space_before = Pt(18)
        h_style.paragraph_format.space_after = Pt(12)
    elif level == 2:
        h_style.font.size = Pt(14)
        h_style.font.bold = True
        h_style.paragraph_format.space_before = Pt(14)
        h_style.paragraph_format.space_after = Pt(10)
    else:
        h_style.font.size = Pt(12)
        h_style.font.bold = True
        h_style.font.italic = True
        h_style.paragraph_format.space_before = Pt(10)
        h_style.paragraph_format.space_after = Pt(8)

# Add title
title = new_doc.add_heading('EFFECT OF A MENTAL HEALTH TRAINING PROGRAMME ON DEPRESSION LITERACY AND REFERRAL COMPETENCE AMONG RELIGIOUS EDUCATORS WORKING WITH ADOLESCENTS', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.size = Pt(16)
    run.font.bold = True

# Add chapters in correct order
chapter_data = [
    ('Ch1', ch1_content),
    ('Ch2', ch2_content),
    ('Ch3', ch3_content),
    ('Ch4', ch4_content),
    ('Ch5', ch5_content),
    ('Ch6', ch6_content),
    ('Ref', ref_content),
    ('App', app_content),
]

for ch_name, content in chapter_data:
    for style_name, text in content:
        text = text.strip()
        if not text:
            new_doc.add_paragraph()
            continue
        
        # Determine style
        if style_name.startswith('Heading 1'):
            p = new_doc.add_heading(text, level=1)
        elif style_name.startswith('Heading 2'):
            p = new_doc.add_heading(text, level=2)
        elif style_name.startswith('Heading 3'):
            p = new_doc.add_heading(text, level=3)
        else:
            p = new_doc.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)

new_doc.save(r'C:\Users\henry\Desktop\GAM APP\Thesis_Final.docx')
print('Final document created successfully')