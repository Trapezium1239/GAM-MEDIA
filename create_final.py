from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Read the backup document
doc = Document(r'C:\Users\henry\Desktop\GAM APP\Thesis_Corrected_backup.docx')

# Extract all content as (style, text) pairs
content = []
for p in doc.paragraphs:
    if p.text.strip() or p.style.name.startswith('Heading'):
        content.append((p.style.name, p.text))

print(f'Total content items: {len(content)}')

# Define the correct order of major sections
section_order = [
    'Chapter One: Introduction',
    'Chapter Two: Literature Review',
    'Chapter Three: Methodology',
    'CHAPTER FOUR',
    'CHAPTER FIVE',
    'CHAPTER SIX',
    'REFERENCE',
    'APPENDIX'
]

# Find start indices for each section
section_starts = {}
for i, (style, text) in enumerate(content):
    for section in section_order:
        if text.startswith(section) or text == section:
            if section not in section_starts:
                section_starts[section] = i

print(f'Section starts: {section_starts}')

# Extract sections in correct order
ordered_content = []
for section in section_order:
    if section in section_starts:
        start = section_starts[section]
        # Find end (next section start or end of content)
        end = len(content)
        for other_section in section_order:
            if other_section in section_starts and section_starts[other_section] > start:
                end = min(end, section_starts[other_section])
        ordered_content.extend(content[start:end])
        print(f'  {section}: {start}-{end} ({end-start} items)')

# Create new document with ordered content
new_doc = Document()

# Set styles
style = new_doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for level in range(1, 4):
    h_style = new_doc.styles[f'Heading {level}']
    h_style.font.name = 'Times New Roman'
    h_style.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        h_style.font.size = Pt(16)
        h_style.font.bold = True
        h_style.paragraph_format.space_before = Pt(18)
        h_style.paragraph_format.space_after = Pt(12)
    elif level == 2:
        h_style.font.size = Pt(14)
        h_style.font.bold = True
        h_style.paragraph_format.space_before = Pt(14)
        h_style.paragraph_format.space_after = Pt(10)
    else:
        h_style.font.size = Pt(12)
        h_style.font.bold = True
        h_style.font.italic = True
        h_style.paragraph_format.space_before = Pt(10)
        h_style.paragraph_format.space_after = Pt(8)

# Add content
for style_name, text in ordered_content:
    text = text.strip()
    if not text:
        new_doc.add_paragraph()
        continue
    
    # Determine style
    if style_name.startswith('Heading 1'):
        p = new_doc.add_heading(text, level=1)
    elif style_name.startswith('Heading 2'):
        p = new_doc.add_heading(text, level=2)
    elif style_name.startswith('Heading 3'):
        p = new_doc.add_heading(text, level=3)
    else:
        p = new_doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

new_doc.save(r'C:\Users\henry\Desktop\GAM APP\Thesis_Final.docx')
print('Final document created successfully')