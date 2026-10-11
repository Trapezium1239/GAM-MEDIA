import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

# Read the fixed text file
with open(r'C:\Users\henry\Desktop\GAM APP\doc_text_fixed.txt', 'r', encoding='utf-8') as f:
    content = f.read()

# Parse paragraphs
paragraphs = []
for line in content.strip().split('\n'):
    match = re.match(r'\[(\d+)\]\s*(.*)', line)
    if match:
        para_num = int(match.group(1))
        text = match.group(2)
        paragraphs.append((para_num, text))

# Create new document
doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Heading styles
for level in range(1, 4):
    h_style = doc.styles[f'Heading {level}']
    h_style.font.name = 'Times New Roman'
    h_style.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        h_style.font.size = Pt(16)
        h_style.font.bold = True
        h_style.paragraph_format.space_before = Pt(18)
        h_style.paragraph_format.space_after = Pt(12)
    elif level == 2:
        h_style.font.size = Pt(14)
        h_style.font.bold = True
        h_style.paragraph_format.space_before = Pt(14)
        h_style.paragraph_format.space_after = Pt(10)
    else:
        h_style.font.size = Pt(12)
        h_style.font.bold = True
        h_style.font.italic = True
        h_style.paragraph_format.space_before = Pt(10)
        h_style.paragraph_format.space_after = Pt(8)

# Helper to add paragraph with style detection
def add_paragraph(doc, text, para_num):
    text = text.strip()
    if not text:
        return
    
    # Detect headings
    if text.startswith('Chapter ') or text == 'Introduction' or text == 'Background':
        if text == 'Chapter One':
            p = doc.add_heading('Chapter One: Introduction', level=1)
        elif text == 'Chapter Two':
            p = doc.add_heading('Chapter Two: Literature Review', level=1)
        elif text == 'Chapter Three':
            p = doc.add_heading('Chapter Three: Methodology', level=1)
        elif text == 'Introduction':
            p = doc.add_heading('Introduction', level=2)
        elif text == 'Background':
            p = doc.add_heading('Background', level=3)
        elif text == 'Literature Review':
            p = doc.add_heading('Literature Review', level=2)
        else:
            p = doc.add_paragraph(text)
    elif re.match(r'^\d+\.\d+\s', text):
        # Sub-section like "1.2 Problem Statement"
        p = doc.add_heading(text, level=2)
    elif re.match(r'^\d+\.\d+\.\d+\s', text):
        p = doc.add_heading(text, level=3)
    elif text.startswith('REFERENCE') or text.startswith('APPENDIX'):
        p = doc.add_heading(text, level=1)
    elif text.startswith('SECTION') or text.startswith('Objective'):
        p = doc.add_heading(text, level=2)
    elif text.startswith('EFFECT OF A MENTAL HEALTH'):
        # Title page
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.name = 'Times New Roman'
    else:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Ensure font consistency
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

# Process paragraphs
for para_num, text in paragraphs:
    # Skip empty
    if not text.strip():
        continue
    add_paragraph(doc, text, para_num)

# Save
doc.save(r'C:\Users\henry\Desktop\GAM APP\Thesis_Corrected.docx')
print('Document saved successfully')